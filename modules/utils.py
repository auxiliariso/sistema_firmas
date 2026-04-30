"""
utils.py — Utilidades de integración con documentos
Funciones para insertar firmas en archivos Excel y Word.
"""

import os
from datetime import datetime


# ─────────────────────────────────────────────
# INTEGRACIÓN CON EXCEL (openpyxl)
# ─────────────────────────────────────────────

def insertar_firma_excel(ruta_archivo: str, firma: dict, hoja: str = None) -> str:
    """
    Inserta la firma digital al final de un archivo Excel existente.
    Si el archivo no existe, lo crea nuevo.

    Parámetros:
        ruta_archivo: ruta completa al .xlsx
        firma:        dict con datos de firma (de firma.py)
        hoja:         nombre de la hoja; None = primera hoja activa

    Retorna la ruta del archivo modificado.
    """
    try:
        from openpyxl import load_workbook, Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        # Cargar o crear el workbook
        if os.path.exists(ruta_archivo):
            wb = load_workbook(ruta_archivo)
            ws = wb[hoja] if hoja and hoja in wb.sheetnames else wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Documento"

        # ── Estilos ──────────────────────────────────
        color_header = "1F3864"   # Azul oscuro corporativo
        color_fila   = "D6E4F0"   # Azul claro

        fuente_titulo  = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        fuente_dato    = Font(name="Calibri", size=10)
        relleno_header = PatternFill("solid", fgColor=color_header)
        relleno_fila   = PatternFill("solid", fgColor=color_fila)
        alineacion     = Alignment(horizontal="left", vertical="center")
        borde_fino = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        # ── Posición: tras la última fila con datos ──
        ultima_fila = ws.max_row + 2  # Dejar una fila de separación

        campos = [
            ("Firmado por", firma['nombre']),
            ("Fecha",       firma['fecha_hora']),
            ("ID Firma",    firma['id_firma']),
            ("Documento",   firma['documento']),
        ]

        # Fila de encabezado de la sección firma
        ws.merge_cells(
            start_row=ultima_fila, start_column=1,
            end_row=ultima_fila, end_column=2
        )
        celda_enc = ws.cell(row=ultima_fila, column=1, value="🔏 FIRMA DIGITAL")
        celda_enc.font      = fuente_titulo
        celda_enc.fill      = relleno_header
        celda_enc.alignment = Alignment(horizontal="center", vertical="center")
        celda_enc.border    = borde_fino
        ws.row_dimensions[ultima_fila].height = 20

        # Filas de datos
        for i, (campo, valor) in enumerate(campos, start=1):
            fila = ultima_fila + i
            c_campo = ws.cell(row=fila, column=1, value=campo)
            c_valor = ws.cell(row=fila, column=2, value=valor)

            for c in (c_campo, c_valor):
                c.font      = fuente_dato
                c.fill      = relleno_fila
                c.alignment = alineacion
                c.border    = borde_fino

            c_campo.font = Font(name="Calibri", bold=True, size=10)
            ws.row_dimensions[fila].height = 16

        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 30

        wb.save(ruta_archivo)
        return ruta_archivo

    except ImportError:
        raise RuntimeError("openpyxl no está instalado. Ejecuta: pip install openpyxl")


# ─────────────────────────────────────────────
# INTEGRACIÓN CON WORD (python-docx)
# ─────────────────────────────────────────────

def insertar_firma_word(ruta_archivo: str, firma: dict) -> str:
    """
    Inserta la firma digital al final de un documento Word existente.
    Si el archivo no existe, lo crea nuevo.

    Parámetros:
        ruta_archivo: ruta completa al .docx
        firma:        dict con datos de firma (de firma.py)

    Retorna la ruta del archivo modificado.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        import copy

        # Cargar o crear el documento
        doc = Document(ruta_archivo) if os.path.exists(ruta_archivo) else Document()

        # ── Separador ────────────────────────────────
        doc.add_paragraph()  # Línea en blanco
        hr = doc.add_paragraph()
        hr.paragraph_format.space_before = Pt(0)
        hr.paragraph_format.space_after  = Pt(0)
        _agregar_linea_horizontal(hr)

        # ── Encabezado de firma ───────────────────────
        p_enc = doc.add_paragraph()
        p_enc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_enc = p_enc.add_run("🔏  FIRMA DIGITAL")
        run_enc.bold      = True
        run_enc.font.size = Pt(11)
        run_enc.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

        # ── Tabla de datos ────────────────────────────
        tabla = doc.add_table(rows=4, cols=2)
        tabla.style = 'Table Grid'

        campos = [
            ("Firmado por", firma['nombre']),
            ("Fecha",       firma['fecha_hora']),
            ("ID Firma",    firma['id_firma']),
            ("Documento",   firma['documento']),
        ]

        for i, (campo, valor) in enumerate(campos):
            fila = tabla.rows[i]

            # Celda etiqueta
            c_etiqueta = fila.cells[0]
            c_etiqueta.text = campo
            run_et = c_etiqueta.paragraphs[0].runs[0]
            run_et.bold           = True
            run_et.font.size      = Pt(10)
            run_et.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

            # Celda valor
            c_valor = fila.cells[1]
            c_valor.text = valor
            run_val = c_valor.paragraphs[0].runs[0]
            run_val.font.size = Pt(10)

        # Anchos de columna (aprox)
        tabla.columns[0].width = Inches(1.5)
        tabla.columns[1].width = Inches(3.5)

        # ── Línea inferior ────────────────────────────
        hr2 = doc.add_paragraph()
        _agregar_linea_horizontal(hr2)

        doc.save(ruta_archivo)
        return ruta_archivo

    except ImportError:
        raise RuntimeError("python-docx no está instalado. Ejecuta: pip install python-docx")


def _agregar_linea_horizontal(paragraph):
    """Agrega una línea horizontal (borde inferior) a un párrafo Word."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '1F3864')
    pBdr.append(bottom)
    pPr.append(pBdr)


# ─────────────────────────────────────────────
# HELPERS GENERALES
# ─────────────────────────────────────────────

def nombre_archivo_seguro(nombre: str) -> str:
    """Elimina caracteres no válidos para nombres de archivo."""
    invalidos = r'\/:*?"<>|'
    for c in invalidos:
        nombre = nombre.replace(c, '_')
    return nombre


def timestamp_archivo() -> str:
    """Retorna timestamp para añadir a nombres de archivo."""
    return datetime.now().strftime("%Y%m%d_%H%M%S")
