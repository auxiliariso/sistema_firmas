"""
firma_celda.py — Inserción de firma en celda específica de Excel
================================================================
La firma contiene únicamente: Nombre del usuario + ID de firma.

Modos de uso:
─────────────────────────────────────────────────────────────
1) Como script directo (interactivo):
       python firma_celda.py

2) Como módulo importado:
       from firma_celda import firmar_celda_excel
       firmar_celda_excel(
           ruta_archivo = "mi_reporte.xlsx",
           celda        = "F10",
           usuario      = "jperez",
           password     = "Firma2026!",
           hoja         = "Hoja1",       # opcional
           modo         = "bloque"       # "bloque" | "linea" | "id"
       )

3) Desde línea de comandos:
       python firma_celda.py --archivo reporte.xlsx --celda F10
                             --usuario jperez --hoja "Hoja1" --modo bloque
─────────────────────────────────────────────────────────────

Modos de inserción:
  bloque  → 2 celdas hacia abajo (Firmado por / ID Firma)   [default]
  linea   → Todo en una sola celda: "Juan Pérez | ID: A82K91LX"
  id      → Solo el ID de firma:    "A82K91LX"
"""

import sys
import os
import argparse
import getpass

_ROOT_DIR = os.path.dirname(os.path.abspath(__file__))
if _ROOT_DIR not in sys.path:
    sys.path.insert(0, _ROOT_DIR)
os.chdir(_ROOT_DIR)

from modules.db    import inicializar_db, verificar_credenciales
from modules.firma import crear_firma


# ═════════════════════════════════════════════
# FUNCIÓN PRINCIPAL EXPORTABLE
# ═════════════════════════════════════════════

def firmar_celda_excel(
    ruta_archivo: str,
    celda:        str,
    usuario:      str,
    password:     str,
    hoja:         str  = None,
    modo:         str  = "bloque",
    crear_si_no_existe: bool = True,
) -> dict:
    """
    Inserta nombre + ID de firma en una celda específica de un archivo Excel.

    Retorna dict con: ok, mensaje, firma, celda, archivo  (o error).
    """

    modos_validos = ("bloque", "linea", "id")
    if modo not in modos_validos:
        return {'ok': False, 'error': f"Modo '{modo}' no válido. Usa: {modos_validos}"}

    # Autenticar
    resultado_auth = verificar_credenciales(usuario, password)
    if resultado_auth is None:
        return {'ok': False, 'error': 'Credenciales incorrectas o usuario no existe.'}
    if isinstance(resultado_auth, dict) and resultado_auth.get('bloqueado'):
        return {'ok': False, 'error': 'Usuario bloqueado. Contacta al administrador.'}

    usuario_data = resultado_auth

    # Generar firma
    nombre_doc = os.path.basename(ruta_archivo)
    firma = crear_firma(usuario_data, nombre_doc)

    # Abrir / crear Excel
    try:
        from openpyxl import load_workbook, Workbook
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    except ImportError:
        return {'ok': False, 'error': 'openpyxl no instalado. Ejecuta: pip install openpyxl'}

    if os.path.exists(ruta_archivo):
        wb = load_workbook(ruta_archivo)
    elif crear_si_no_existe:
        wb = Workbook()
    else:
        return {'ok': False, 'error': f"Archivo no encontrado: {ruta_archivo}"}

    ws = wb[hoja] if (hoja and hoja in wb.sheetnames) else wb.active

    # Parsear celda
    celda = celda.upper().strip()
    try:
        col_letra, fila_num = coordinate_from_string(celda)
        col_num = column_index_from_string(col_letra)
    except Exception:
        return {'ok': False, 'error': f"Celda '{celda}' no válida. Usa formato Excel: A1, B3, F10…"}

    # Insertar
    if modo == "bloque":
        _insertar_bloque(ws, fila_num, col_num, firma)
    elif modo == "linea":
        _insertar_linea(ws, fila_num, col_num, firma)
    elif modo == "id":
        _insertar_solo_id(ws, fila_num, col_num, firma)

    wb.save(ruta_archivo)

    return {
        'ok':      True,
        'mensaje': f"Firma insertada en {ws.title}!{celda} (modo: {modo})",
        'firma': {
            'nombre':   firma['nombre'],
            'id_firma': firma['id_firma'],
            'documento': firma['documento'],
        },
        'celda':   f"{ws.title}!{celda}",
        'archivo': ruta_archivo,
    }


# ═════════════════════════════════════════════
# ESTILOS COMPARTIDOS
# ═════════════════════════════════════════════

def _estilos():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    return {
        'fuente_header': Font(name="Calibri", bold=True, color="FFFFFF", size=10),
        'fuente_label':  Font(name="Calibri", bold=True, size=10, color="1F3864"),
        'fuente_valor':  Font(name="Calibri", size=10),
        'relleno_header': PatternFill("solid", fgColor="1F3864"),
        'relleno_fila':   PatternFill("solid", fgColor="D6E4F0"),
        'alineacion':     Alignment(horizontal="left", vertical="center", wrap_text=True),
        'borde': Border(
            left=Side(style='thin'),  right=Side(style='thin'),
            top=Side(style='thin'),   bottom=Side(style='thin')
        ),
    }


# ═════════════════════════════════════════════
# MODOS DE INSERCIÓN
# ═════════════════════════════════════════════

def _insertar_bloque(ws, fila: int, col: int, firma: dict):
    """
    Modo BLOQUE — 3 filas:
        [Encabezado] 🔏 FIRMA DIGITAL  (merge 2 cols)
        [F+1]  Firmado por  │  Juan Pérez
        [F+2]  ID Firma     │  A82K91LX
    """
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    s = _estilos()

    # Encabezado
    try:
        ws.merge_cells(start_row=fila, start_column=col,
                       end_row=fila,   end_column=col + 1)
    except Exception:
        pass

    c_enc = ws.cell(row=fila, column=col, value="🔏 FIRMA DIGITAL")
    c_enc.font      = s['fuente_header']
    c_enc.fill      = s['relleno_header']
    c_enc.alignment = Alignment(horizontal="center", vertical="center")
    c_enc.border    = s['borde']
    ws.row_dimensions[fila].height = 18

    # Solo nombre + ID (sin fecha)
    datos = [
        ("Firmado por", firma['nombre']),
        ("ID Firma",    firma['id_firma']),
    ]

    for offset, (etiqueta, valor) in enumerate(datos, start=1):
        f = fila + offset
        c_lbl = ws.cell(row=f, column=col,     value=etiqueta)
        c_val = ws.cell(row=f, column=col + 1, value=valor)

        c_lbl.font      = s['fuente_label']
        c_lbl.fill      = s['relleno_fila']
        c_lbl.alignment = s['alineacion']
        c_lbl.border    = s['borde']

        c_val.font      = s['fuente_valor']
        c_val.fill      = s['relleno_fila']
        c_val.alignment = s['alineacion']
        c_val.border    = s['borde']

        ws.row_dimensions[f].height = 15

    col_a = get_column_letter(col)
    col_b = get_column_letter(col + 1)
    if ws.column_dimensions[col_a].width < 14:
        ws.column_dimensions[col_a].width = 14
    if ws.column_dimensions[col_b].width < 26:
        ws.column_dimensions[col_b].width = 26


def _insertar_linea(ws, fila: int, col: int, firma: dict):
    """
    Modo LINEA — una sola celda:
        Juan Pérez  |  ID: A82K91LX
    """
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    s = _estilos()
    texto = f"{firma['nombre']}  |  ID: {firma['id_firma']}"

    c = ws.cell(row=fila, column=col, value=texto)
    c.font      = s['fuente_valor']
    c.fill      = s['relleno_fila']
    c.border    = s['borde']
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[fila].height = 16

    col_letra = get_column_letter(col)
    ancho = len(texto) * 0.9
    if ws.column_dimensions[col_letra].width < ancho:
        ws.column_dimensions[col_letra].width = min(ancho, 80)


def _insertar_solo_id(ws, fila: int, col: int, firma: dict):
    """
    Modo ID — solo el ID de firma:
        A82K91LX
    """
    from openpyxl.styles import Alignment, Font

    c = ws.cell(row=fila, column=col, value=firma['id_firma'])
    c.font      = Font(name="Courier New", bold=True, size=10, color="1F3864")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border    = _estilos()['borde']
    ws.row_dimensions[fila].height = 15


# ═════════════════════════════════════════════
# MODO INTERACTIVO
# ═════════════════════════════════════════════

def _modo_interactivo():
    print("\n" + "═" * 52)
    print("  🔏  Firma Digital — Celda Específica en Excel")
    print("═" * 52)

    print("\n📂 Ruta del archivo Excel (.xlsx):")
    ruta = input("   → ").strip().strip('"')
    if not ruta:
        print("❌ Ruta vacía. Saliendo.")
        return

    print("\n📋 Nombre de la hoja (Enter = hoja activa):")
    hoja = input("   → ").strip() or None

    print("\n📌 Celda de destino (ej: F10, B3, A1):")
    celda = input("   → ").strip()
    if not celda:
        print("❌ Celda vacía. Saliendo.")
        return

    print("\n🎨 Modo de inserción:")
    print("   1  bloque  → encabezado + nombre + ID  [recomendado]")
    print("   2  linea   → todo en 1 celda")
    print("   3  id      → solo el ID de firma")
    opcion = input("   → ").strip()
    modos_map = {'1': 'bloque', '2': 'linea', '3': 'id',
                 'bloque': 'bloque', 'linea': 'linea', 'id': 'id'}
    modo = modos_map.get(opcion, 'bloque')
    print(f"   Modo: {modo}")

    print("\n🔐 Credenciales:")
    usuario  = input("   Usuario  → ").strip()
    password = getpass.getpass("   Contraseña → ")

    print("\n⏳ Procesando...")
    resultado = firmar_celda_excel(
        ruta_archivo=ruta, celda=celda,
        usuario=usuario, password=password,
        hoja=hoja, modo=modo,
    )

    print()
    if resultado['ok']:
        f = resultado['firma']
        print("✅ FIRMA INSERTADA")
        print(f"   Archivo:     {resultado['archivo']}")
        print(f"   Celda:       {resultado['celda']}")
        print(f"   Firmado por: {f['nombre']}")
        print(f"   ID Firma:    {f['id_firma']}")
    else:
        print(f"❌ ERROR: {resultado['error']}")

    input("\nPresiona Enter para salir...")


# ═════════════════════════════════════════════
# LÍNEA DE COMANDOS
# ═════════════════════════════════════════════

def _modo_cli():
    parser = argparse.ArgumentParser(
        description='Inserta nombre + ID de firma en una celda específica de Excel.',
        formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument('--archivo',  required=True,  help='Ruta al archivo .xlsx')
    parser.add_argument('--celda',    required=True,  help='Celda destino (ej: F10, B3)')
    parser.add_argument('--usuario',  required=True,  help='Nombre de usuario')
    parser.add_argument('--hoja',     default=None,   help='Nombre de hoja (default: activa)')
    parser.add_argument('--modo',     default='bloque',
                        choices=['bloque', 'linea', 'id'],
                        help='Modo de inserción (default: bloque)')
    args = parser.parse_args()

    password = getpass.getpass(f"Contraseña para '{args.usuario}': ")

    resultado = firmar_celda_excel(
        ruta_archivo=args.archivo, celda=args.celda,
        usuario=args.usuario, password=password,
        hoja=args.hoja, modo=args.modo,
    )

    if resultado['ok']:
        print(f"✅ {resultado['mensaje']}")
        print(f"   ID Firma: {resultado['firma']['id_firma']}")
        sys.exit(0)
    else:
        print(f"❌ Error: {resultado['error']}")
        sys.exit(1)


# ═════════════════════════════════════════════
if __name__ == '__main__':
    inicializar_db()
    if len(sys.argv) > 1:
        _modo_cli()
    else:
        _modo_interactivo()