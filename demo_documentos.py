"""
demo_documentos.py — Demostración de integración con documentos
Genera archivos Excel y Word de ejemplo con firmas insertadas.
Ejecutar directamente: python demo_documentos.py
"""

import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from modules.db     import inicializar_db, crear_usuario
from modules.firma  import crear_firma, formatear_firma_texto
from modules.utils  import insertar_firma_excel, insertar_firma_word


def demo_excel(firma: dict, carpeta_salida: str):
    """Crea un Excel de ejemplo con datos ficticios y firma al final."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    ruta = os.path.join(carpeta_salida, "reporte_ventas_firmado.xlsx")

    # Crear Excel con datos de ejemplo
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"

    # Encabezado
    ws['A1'] = "Reporte de Ventas Q1 2026"
    ws['A1'].font = Font(bold=True, size=14, color="1F3864")

    # Cabeceras de tabla
    encabezados = ["Producto", "Unidades", "Precio Unitario", "Total"]
    for col, enc in enumerate(encabezados, start=1):
        celda = ws.cell(row=3, column=col, value=enc)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3864")
        celda.alignment = Alignment(horizontal="center")

    # Datos de ejemplo
    datos = [
        ("Laptop Pro",    15, 25000, 375000),
        ("Mouse Inalám.", 50,   450,  22500),
        ("Teclado Mec.",  30,  1200,  36000),
        ("Monitor 27\"",  10, 12000, 120000),
    ]
    for fila_i, fila in enumerate(datos, start=4):
        for col_i, valor in enumerate(fila, start=1):
            ws.cell(row=fila_i, column=col_i, value=valor)

    # Total general
    ws.cell(row=8, column=3, value="TOTAL:").font = Font(bold=True)
    ws.cell(row=8, column=4, value=553500).font = Font(bold=True, color="1E8449")

    # Ajustar anchos
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 14

    wb.save(ruta)

    # Ahora insertar la firma
    insertar_firma_excel(ruta, firma)
    print(f"  ✓ Excel creado con firma: {ruta}")
    return ruta


def demo_word(firma: dict, carpeta_salida: str):
    """Crea un documento Word de ejemplo con firma al final."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    ruta = os.path.join(carpeta_salida, "contrato_servicio_firmado.docx")

    # Crear documento de ejemplo
    doc = Document()

    # Título
    titulo = doc.add_heading("CONTRATO DE PRESTACIÓN DE SERVICIOS", level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    doc.add_paragraph()

    # Cuerpo del contrato (texto de ejemplo)
    parrafos = [
        "En la Ciudad de México, a 30 de abril de 2026, las partes acuerdan los siguientes términos:",
        "",
        "PRIMERA. — El Prestador se compromete a entregar los servicios de consultoría tecnológica "
        "en los plazos establecidos en el Anexo A del presente contrato.",
        "",
        "SEGUNDA. — El Cliente pagará la contraprestación acordada de $120,000.00 MXN en tres "
        "parcialidades iguales, conforme al calendario de pagos del Anexo B.",
        "",
        "TERCERA. — Ambas partes acuerdan que cualquier disputa será resuelta mediante arbitraje "
        "en términos de la legislación mercantil aplicable.",
        "",
        "Leído el presente instrumento y enteradas las partes de su contenido y alcance legal, "
        "lo firman al calce.",
    ]
    for p in parrafos:
        doc.add_paragraph(p)

    doc.save(ruta)

    # Insertar la firma digital
    insertar_firma_word(ruta, firma)
    print(f"  ✓ Word creado con firma:  {ruta}")
    return ruta


def main():
    print("=" * 50)
    print("  Demo de integración con documentos")
    print("=" * 50)

    # Inicializar DB y crear usuario demo
    inicializar_db()
    crear_usuario("demo_user", "DemoPass123!", "Ana Rodríguez")

    usuario_data = {
        'usuario': 'demo_user',
        'nombre':  'Ana Rodríguez'
    }

    # Crear firma
    print("\n[FIRMA] Generando firma digital...")
    firma = crear_firma(usuario_data, "documentos_demo")
    print(f"\n{formatear_firma_texto(firma)}\n")

    # Carpeta de salida
    carpeta = os.path.join(os.path.dirname(__file__), "output_examples")
    os.makedirs(carpeta, exist_ok=True)

    print("[DOC] Generando documentos de ejemplo...")
    demo_excel(firma, carpeta)
    demo_word(firma, carpeta)

    print(f"\n✅ Archivos guardados en: {carpeta}")
    print("\nPuedes abrir los archivos para ver la firma insertada.")


if __name__ == "__main__":
    main()
